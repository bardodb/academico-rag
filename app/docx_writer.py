"""Markdown do trabalho para .docx no padrão da NBR 14724."""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Arial"
HEADING_LEVEL = {"##": 1, "###": 2, "####": 3, "#####": 4, "######": 5}
LIST_ITEM = re.compile(r"^\s*(?:[-*•]|(\d+)[.)])\s+(.*)$")
TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def save_docx(markdown: str, path: Path) -> None:
    doc = Document()
    _setup(doc)
    lines = markdown.splitlines()
    title = lines[0].lstrip("# ").strip() if lines and lines[0].startswith("# ") else "Trabalho acadêmico"
    _cover(doc, title)
    _toc(doc)
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    _page_number(body)
    _body(doc, lines[1:] if lines and lines[0].startswith("# ") else lines)
    doc.save(path)


def _body(doc: Document, lines: list[str]) -> None:
    i, first_primary, in_gaps = 0, True, False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            _mono(doc, block)
            i += 1
            continue
        heading = re.match(r"^(#{2,6})\s+(.*)$", stripped)
        if heading:
            level = HEADING_LEVEL.get(heading.group(1), 4)
            text = heading.group(2).replace("**", "").strip()
            h = doc.add_heading(text, level=level)
            if level == 1:
                h.paragraph_format.page_break_before = not first_primary
                first_primary = False
                in_gaps = text.lower().startswith("relatório de lacunas")
                if text.upper() in {"REFERÊNCIAS", "RELATÓRIO DE LACUNAS"}:
                    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not TABLE_SEPARATOR.match(lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _table(doc, rows)
            continue
        item = LIST_ITEM.match(line)
        if item:
            marker = f"{item.group(1)}." if item.group(1) else "•"
            if _in_references(doc) and not in_gaps:
                _reference(doc, item.group(2))
            else:
                _item(doc, item.group(2), marker)
            i += 1
            continue
        _paragraph(doc, stripped)
        i += 1


def _in_references(doc: Document) -> bool:
    for paragraph in reversed(doc.paragraphs):
        if paragraph.style.name == "Heading 1":
            return paragraph.text.strip().upper() == "REFERÊNCIAS"
    return False


def _runs(paragraph, text: str, size: float | None = None) -> None:
    text = text.replace("`", "")
    for index, part in enumerate(re.split(r"\*\*", text)):
        for jndex, piece in enumerate(re.split(r"(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)", part)):
            if not piece:
                continue
            run = paragraph.add_run(piece)
            run.bold = index % 2 == 1
            run.italic = jndex % 2 == 1
            if size:
                run.font.size = Pt(size)


def _paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    _runs(p, text)


def _item(doc: Document, text: str, marker: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    _runs(p, f"{marker}\u00a0{text}")


def _reference(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    _runs(p, text)


def _table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for c_index in range(cols):
            value = row[c_index] if c_index < len(row) else ""
            cell = cells[c_index]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
            _runs(p, value.replace("<br>", " "), 9)
            if r_index == 0:
                for run in p.runs:
                    run.bold = True
                _shade(cell)
    _source(doc)


def _mono(doc: Document, block: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.text = ""
    for index, text in enumerate(block or [" "]):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run(text or " ")
        run.font.name = "Courier New"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(8.5)
    _source(doc)


def _source(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    bold = p.add_run("Fonte: ")
    bold.bold = True
    bold.font.size = Pt(10)
    p.add_run("Elaborado pelo autor.").font.size = Pt(10)


def _shade(cell, color: str = "D9D9D9") -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._element.get_or_add_tcPr().append(shd)


def _cover(doc: Document, title: str) -> None:
    def center(text: str, bold: bool = False, before: int = 0) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        run = p.add_run(text)
        run.bold = bold

    center("NOME DA INSTITUIÇÃO", bold=True)
    center("NOME DO CURSO", bold=True)
    center("NOME COMPLETO DO(A) ESTUDANTE", bold=True, before=140)
    center(title.upper(), bold=True, before=140)
    center("Cidade", before=200)
    center("Ano")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _toc(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("SUMÁRIO").bold = True
    p = doc.add_paragraph()
    _field(p, 'TOC \\o "1-5" \\h \\z \\u', "Clique com o botão direito e escolha “Atualizar campo”.")


def _page_number(section) -> None:
    section.header.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _field(p, "PAGE", "")


def _field(paragraph, instruction: str, placeholder: str) -> None:
    def fld(kind: str) -> None:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(el)

    fld("begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    paragraph.add_run()._r.append(instr)
    fld("separate")
    if placeholder:
        paragraph.add_run(placeholder)
    fld("end")


def _setup(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.left_margin = Cm(3), Cm(3)
    section.bottom_margin, section.right_margin = Cm(2), Cm(2)
    _style(doc.styles["Normal"], 12, False)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing, normal.space_after, normal.space_before = 1.5, Pt(0), Pt(0)
    for level, caps, bold in ((1, True, True), (2, True, False), (3, False, True), (4, False, False), (5, False, False)):
        style = doc.styles[f"Heading {level}"]
        _style(style, 12, bold)
        style.font.all_caps = caps
        pf = style.paragraph_format
        pf.space_before, pf.space_after = Pt(0 if level == 1 else 12), Pt(12)
        pf.line_spacing, pf.keep_with_next, pf.alignment = 1.5, True, WD_ALIGN_PARAGRAPH.LEFT
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)


def _style(style, size: int, bold: bool) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
